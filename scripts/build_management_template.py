#!/usr/bin/env python3
"""Build a blank 12-page management-interpretation template. No real report is read.
Requires python-pptx. This is a template builder, not a report generation engine.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement
from build_templates import (group, slide_new, header, text, rect, ln, radar, block, notes,
                             clean_package, DIMS, W, H, BLUE, RED, GOLD, INK, MUTED, PALE,
                             ROSE, LIGHT, LINE, WHITE, PP_ALIGN)


def clear(slide):
    for shape in list(slide.shapes):
        el = shape._element
        el.getparent().remove(el)


def rewrite_text(shape, value):
    # Keep the first run's formatting; useful for changing page labels in unchanged layouts.
    first = shape.text_frame.paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        for run in list(first.runs)[1:]:
            first._p.remove(run._r)
        for paragraph in list(shape.text_frame.paragraphs)[1:]:
            shape.text_frame._txBody.remove(paragraph._p)
    else:
        shape.text = value


def difference_card(s, prefix, x, y, w, title, body, evidence, color=BLUE):
    rect(s,prefix+'.bg',x,y,w,1.08,PALE if color==BLUE else ROSE,rounded=True)
    text(s,prefix+'.title',title,x+.13,y+.07,w-.26,.29,12.3,color,True)
    text(s,prefix+'.body',body,x+.13,y+.40,w-.26,.37,10.7)
    text(s,prefix+'.evidence',evidence,x+.13,y+.84,w-.26,.16,8.7,MUTED)


def build_deck():
    p = group()  # Constructs blank shapes only, not a copy of anyone's report.
    original = list(p.slides)
    # Page 3: state the point-label requirement in the actual placeholder.
    for sh in original[2].shapes:
        if sh.has_text_frame and sh.name.endswith('.replace'):
            rewrite_text(sh,'[均值 / 中位值折线；数值直接标在点位旁]')
    notes(original[2], '按当批数据生成均值和中位值曲线，数值标在对应点旁，颜色和线型与图例一致。不要改成横轴下面的数据表。无实际数据，不能把占位图当统计结果。')

    # 4: group + individual variation on one page.
    s=original[3];clear(s)
    header(s,4,'HPI：{群体做事特点，以及需要区别看待的个体}','日常工作风格  /  群体共性 + 有依据的高低差异；不作能力排名')
    radar(s,'m4.radar',.56,1.54,5.49,3.94,DIMS['HPI'],'本批HPI群体雷达图\n均值 / 中位值 / 个体轮廓')
    for j in range(3):
        block(s,f'm4.finding{j}',6.26,1.55+j*1.25,6.48,1.12,
              f'[群体结论{j+1}：具体怎样做事]',
              '[在任务中的可能表现；对理解、推进或合作的影响。]',
              '依据：{本批分布或关键统计}')
    text(s,'m4.diff.heading','差异个体：有依据时兼顾高于与低于群体的不同做法',.68,5.47,12,.27,12,INK,True)
    for j in range(3):
        difference_card(s,f'm4.diff{j}',.65+j*4.19,5.88,4.00,
          '{显示名/编号}｜{与群体不同在哪里}',
          '[可能的做法与工作影响；\n不能直接套用群体结论。]',
          '依据：{个人/群体；必要时补HPI组合}')
    notes(s,'同页左侧群体雷达图、右侧群体结论、底部差异卡。差异须基于本批原分与解释价值，高低两端有证据才选；不固定三人、不把相对低分叫低能力或统计显著。姓名模式依本次受众授权，默认代号。')

    # 6: focus on main HDS scales, subscale detail has its own page.
    for sh in original[5].shapes:
        if sh.has_text_frame and sh.name.endswith('.subscale'):
            rewrite_text(sh,'依据：{主量表分布 / 实际来源}')
    notes(original[5],'本页只解释HDS主量表与压力条件；子维度拆在下一页。重点依据本批分布选择，不固定四项。风险不是已发生行为，说明具体工作影响。')

    # 8: heatmap with configurable identity mode, four-band key and observation cards.
    s=original[6];clear(s)
    header(s,8,'HDS：{压力反应组合不同，工作关注需要区分}','热力图  /  默认编号；实名仅在用户指定、受众授权、映射确认后使用')
    text(s,'m8.summary','[先概括本批不同压力组合对协作与交付的主要影响]',.69,1.54,12,.35,15,INK,True)
    x0,y0=.69,2.26;ww,hh=7.31,3.91;cw=ww/12;rh=hh/7
    labels=['显示名']+DIMS['HDS']
    for rr in range(7):
        for cc in range(12):
            xx=x0+cc*cw;yy=y0+rr*rh
            rect(s,f'm8.cell{rr}.{cc}.bg',xx,yy,cw-.018,rh-.018,'EDF3FC' if rr==0 else LIGHT,LINE)
            label=labels[cc] if rr==0 else (f'成员{rr:02}' if cc==0 else '—')
            text(s,f'm8.cell{rr}.{cc}.text',label,xx+.009,yy+.035,cw-.036,rh-.07,8.5,BLUE if cc==0 else INK,rr==0,PP_ALIGN.CENTER)
    text(s,'m8.identity.note','显示模式：{编号 / 授权缩写 / 授权实名}',.74,1.98,7.0,.24,10.3,MUTED)
    text(s,'m8.key.title','图例',8.16,1.98,.73,.25,11,INK,True)
    for i,(c,lab) in enumerate([('EF5B64','{高档}'),('F0A04B','{中高档}'),('D5BE47','{中低档}'),('56AA86','{低档}')]):
        yy=2.40+i*.88
        rect(s,f'm8.key{i}.color',8.16,yy,.16,.66,c)
        text(s,f'm8.key{i}.label',lab,8.39,yy,.58,.29,9.8,MUTED)
        text(s,f'm8.key{i}.range','{区间}',8.39,yy+.35,.58,.22,8.6,MUTED)
    text(s,'m8.observer.heading','观察对象 / 共性组合',9.24,1.98,3.38,.26,12.5,BLUE,True)
    for j in range(3):
        difference_card(s,f'm8.observer{j}',9.20,2.39+j*1.25,3.49,
                        '{对象/组合}｜{重点}',
                        '[压力下可能怎样做，\n影响哪一步工作。]',
                        '依据：{分值/人数；选择理由}')
    text(s,'m8.relabel','标签说明：{原报告区间及称谓；如按参考改显示标签，写清映射，阈值不变。}',.74,6.49,11.88,.37,10.1,MUTED)
    notes(s,'六行与三张观察卡只是模板容量，不是本批人数或筛选规则。按N增减或分页；保留完整表。色带为显示占位，须核对源区间；显示标签变化不得改变阈值。观察对象按本批证据选，不默认五人或某个固定分数规则。真实身份映射不发布。')

    # 9: common motives + differentiated work arrangements.
    s=original[7];clear(s)
    header(s,9,'MVPI：{共同投入来源，以及需要区别沟通的偏好}','动机与价值偏好  /  任务沟通、工作环境与反馈；不等同业务能力')
    radar(s,'m9.radar',.56,1.54,5.49,3.94,DIMS['MVPI'],'本批MVPI群体雷达图\n均值 / 中位值 / 个体轮廓')
    for j in range(3):
        block(s,f'm9.finding{j}',6.26,1.55+j*1.25,6.48,1.12,
              f'[共同驱动{j+1}：更重视什么]',
              '[说明哪类任务沟通或安排可能更能激发投入。]',
              '依据：{本批分布；成效仍需确认}')
    text(s,'m9.diff.heading','差异个体：同一种任务安排或认可方式，不一定对所有人都有效',.68,5.47,12,.27,12,INK,True)
    for j in range(3):
        difference_card(s,f'm9.diff{j}',.65+j*4.19,5.88,4.00,
          '{显示名/编号}｜{偏好区别}',
          '[怎样沟通任务、给予责任或反馈；\n说明需要与本人确认。]',
          '依据：{相对分值或相关量表组合}')
    notes(s,'共同驱动与差异个体同时呈现。MVPI不是能力或道德判断，不声称最有效激励。低传统不等于不守规则。工作安排为待验证建议，引用必须来自当批报告。')

    # 7: additional subscale slide.
    sub=slide_new(p);header(sub,7,'HDS子维度：{具体突出什么，会怎样影响工作}','原图统计 → 可能的工作表现 → 协作与交付影响；不从主量表猜格数')
    xs=[.68,4.18,8.28];ws=[3.34,3.94,4.38]
    for j,label in enumerate(['具体子维度与人数','可能的工作表现','对协作与交付的影响']):
        rect(sub,f'm7.head{j}.bg',xs[j],1.63,ws[j],.50,PALE if j!=2 else ROSE)
        text(sub,f'm7.head{j}.text',label,xs[j]+.15,1.72,ws[j]-.30,.30,13.3,BLUE if j!=2 else RED,True)
    for i in range(4):
        yy=2.33+i*1.00
        values=[f'{{主量表 {i+1}}}\n{{突出子维度 / 人数 / 有效N}}',
                '[什么条件下，可能怎样反应；\n说明观察行为，不给人贴标签。]',
                '[有依据时说明正向作用；\n再说明使用过度时的一步工作后果。]']
        for j,val in enumerate(values):
            rect(sub,f'm7.row{i}.{j}.bg',xs[j],yy,ws[j],.86,LIGHT if i%2==0 else WHITE,LINE)
            text(sub,f'm7.row{i}.{j}.text',val,xs[j]+.14,yy+.11,ws[j]-.28,.62,11.4,INK,j==0)
    text(sub,'m7.limit','{原图、来源页码、格数定义与计数规则}；“0人满格”不等于没有倾向；缺页写有效N。',.74,6.65,11.85,.29,10.1,MUTED)
    notes(sub,'只有可读原图及定义时填写。满格计数是项目规则，不等同主量表风险概率。不得强行把风险美化成优点，未提供资料时注明未统计或按请求删除本页。')

    # 10: work observation matrix.
    work=slide_new(p);header(work,10,'工作观察：{哪些表现会影响判断、推进、交付和协作}','在真实任务中验证  /  场景、表现与影响、观察行为、支持方式相互对应')
    xs=[.67,2.64,6.80,9.86];ws=[1.80,3.99,2.89,2.80]
    for j,label in enumerate(['工作场景','可能的表现与工作影响','应观察的行为','对应支持方式']):
        rect(work,f'm10.head{j}.bg',xs[j],1.64,ws[j],.52,PALE)
        text(work,f'm10.head{j}.text',label,xs[j]+.10,1.72,ws[j]-.20,.35,12.3,BLUE,True)
    scenes=['[方案评审]','[需求与排期]','[关键交付]','[跨团队协作]']
    for i in range(4):
        yy=2.38+i*.98
        vals=[scenes[i],'[根据本批证据写可能的做法；\n再说明对哪一步工作有帮助或代价。]',
              '[可看见、可核对的动作；\n不是笼统能力标签。]',
              '[与前面行为对应的\n任务安排或反馈动作。]']
        for j,val in enumerate(vals):
            rect(work,f'm10.row{i}.{j}.bg',xs[j],yy,ws[j],.86,LIGHT if i%2==0 else WHITE,LINE)
            text(work,f'm10.row{i}.{j}.text',val,xs[j]+.10,yy+.11,ws[j]-.20,.62,11.2,INK,j==0)
    text(work,'m10.limit','以上场景须按用户工作背景选择；没有实际记录时一律作为待验证假设，不描述成已发生问题。',.74,6.65,11.85,.31,10.2,MUTED)
    notes(work,'四个场景仅示范结构。只要结论时，可删除支持方式列，不自动加培养方案。建议不能超出前述证据；实际任务记录与测评推测要分开。')

    # 11: three columns and a bottom strip.
    summary=slide_new(p);header(summary,11,'群体总结：{最值得保留的优势与发展的重点}','可放大的优势 / 需管理的风险 / 对应的发展建议')
    rect(summary,'m11.overview.bg',.68,1.55,11.98,.66,PALE,rounded=True)
    text(summary,'m11.overview','[一句完整的群体画像；再概括当前最重要的工作关注，不只列分数。]',.86,1.70,11.62,.35,15,INK,True)
    headings=['可放大的优势','需管理的风险','对应的发展建议']
    subheads=['日常风格 × 投入动机','高压与责任扩大时观察','从实际任务与反馈中发展']
    for col in range(3):
        x=.68+col*4.07;c=RED if col==1 else BLUE
        rect(summary,f'm11.col{col}.bg',x,2.51,3.84,3.70,LIGHT,rounded=True)
        rect(summary,f'm11.col{col}.stripe',x,2.48,3.84,.045,c)
        text(summary,f'm11.col{col}.heading',headings[col],x+.17,2.67,3.50,.34,17,c,True)
        text(summary,f'm11.col{col}.sub',subheads[col],x+.17,3.07,3.50,.27,9.8,MUTED)
        for row in range(3):
            yy=3.51+row*.83
            title=['[优势：帮助哪一步工作]','[风险：影响什么工作]','[建议：对应前面的行为]'][col]
            body=['[说明任务中的可能价值；\n成效仍须结合实际成果验证。]',
                  '[说明具体行为与一步后果；\n不是已发生问题的断言。]',
                  '[写具体安排、反馈或观察；\n不新增无依据的培养项目。]'][col]
            text(summary,f'm11.col{col}.row{row}.title',title,x+.17,yy,3.50,.25,11.8,INK,True)
            text(summary,f'm11.col{col}.row{row}.body',body,x+.17,yy+.30,3.50,.39,10.5)
    rect(summary,'m11.impact.bg',.68,6.46,11.98,.49,WHITE,LINE)
    text(summary,'m11.impact.title','业务与组织影响',.83,6.55,1.44,.23,10.5,INK,True)
    text(summary,'m11.impact.positive','正向价值：[对哪些任务有帮助]',2.43,6.55,4.78,.23,10.5,BLUE)
    text(summary,'m11.impact.risk','需要防范：[什么条件下影响交付]',7.43,6.55,5.06,.23,10.5,RED)
    notes(summary,'三栏回应前文，底部说明业务影响。只要结论时第三栏替换为观察重点，不自动给发展建议。不要把旧群体结论和数据迁入新批次。')

    # Reorder; drop only the unused blank legacy observation slide.
    ids=list(p.slides._sldIdLst)
    order=[0,1,2,3,4,5,10,6,7,11,12,9]
    p.part.drop_rel(ids[8].rId)
    for el in list(p.slides._sldIdLst):
        p.slides._sldIdLst.remove(el)
    for idx in order:
        p.slides._sldIdLst.append(ids[idx])
    guide=p.slides[11]
    for sh in guide.shapes:
        if sh.name=='p10.index':rewrite_text(sh,'12')
        if sh.has_text_frame and '身份映射' in sh.text:
            rewrite_text(sh,sh.text.replace('群体报告默认编号，身份映射不进共享备注。','群体默认编号；实名或映射附录需本次受众授权。'))
    p.core_properties.title='Hogan群体管理解读空白模板（12页）'
    return p


def build(out: Path) -> dict:
    out.mkdir(parents=True,exist_ok=True)
    target=out/'group-management-template.pptx'
    mapping=out/'management-template-map.json'
    if target.exists() or mapping.exists():
        raise ValueError('目标已有同名文件，请使用新的空白目录')
    p=build_deck()
    for sl in p.slides:
        for sh in sl.shapes:
            for el in sh._element.xpath('.//a:effectRef'):el.set('idx','0')
            for sp in sh._element.xpath('./p:spPr'):
                for old in sp.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst'):sp.remove(old)
                sp.append(OxmlElement('a:effectLst'))
    p.save(target);clean_package(target)
    result={'filename':target.name,'profile':'management12','size_inches':[W,H],'page_count':len(p.slides),'pages':[]}
    for i,sl in enumerate(p.slides,1):
        fields=[]
        for sh in sl.shapes:
            if sh.has_text_frame:
                fields.append({'name':sh.name,'shape_id':sh.shape_id,'x':round(sh.left/914400,4),
                               'y':round(sh.top/914400,4),'w':round(sh.width/914400,4),
                               'h':round(sh.height/914400,4),'placeholder':sh.text})
        result['pages'].append({'page':i,'fields':fields})
    result['note']='从空白图形构造；人数、标色、标签和身份显示须按本次任务确认。只重建模板，不填写真实分析。'
    mapping.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return result
if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();build(args.out);print('12页空白模板及字段映射已创建；仍需渲染审查。')
