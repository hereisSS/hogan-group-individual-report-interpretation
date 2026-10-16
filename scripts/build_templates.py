#!/usr/bin/env python3
"""Rebuild blank, editable templates without personal data or inherited Office assets.
Requires python-pptx. This builds templates, NOT an analysis/report generator.
Run from the repository root. Does not read any real report.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import json
import math
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

W,H=13.333333,7.5
FONT='Noto Sans CJK SC'
INK='263345';MUTED='6F7B8D';BLUE='3078F5';RED='E8565B';GOLD='B48B28';LINE='DFE5EF'
PALE='F3F7FD';ROSE='FCF3F2';LIGHT='FAFBFD';WHITE='FFFFFF'
DIMS={'HPI':['调适','抱负','社交','人际敏感度','审慎','好奇','学习方式'],
      'HDS':['激动','多疑','谨慎','内敛','消极','自大','狡猾','戏剧化','幻想','苛求','恭顺'],
      'MVPI':['认可','权力','享乐','利他','归属','传统','保障','商业','美感','科学']}


def rgb(c):return RGBColor.from_string(c)
def text(slide,name,value,x,y,w,h,size=13,color=INK,bold=False,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=name
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.015);tf.margin_top=tf.margin_bottom=Inches(.006)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=1.12
        r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
    return sh

def rect(s,name,x,y,w,h,fill=WHITE,stroke=None,rounded=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.name=name
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if stroke:sh.line.color.rgb=rgb(stroke);sh.line.width=Pt(.65)
    else:sh.line.fill.background()
    if rounded:
        try:sh.adjustments[0]=.035
        except Exception:pass
    return sh

def ln(s,name,x1,y1,x2,y2,color=LINE,width=.7,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.name=name
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    if dash:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def ring(s,name,x,y,d,color=LINE):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.name=name
    sh.fill.background();sh.line.color.rgb=rgb(color);sh.line.width=Pt(.7)
    return sh

def prs_new(title):
    p=Presentation();p.slide_width=Inches(W);p.slide_height=Inches(H)
    p.core_properties.author='Hogan Workflow Template';p.core_properties.last_modified_by='Hogan Workflow Template'
    p.core_properties.title=title;p.core_properties.subject='Blank workflow template; no actual assessment data'
    p.core_properties.comments='';p.core_properties.keywords='blank template'
    p.core_properties.created=datetime(2026,10,15,tzinfo=timezone.utc);p.core_properties.modified=datetime(2026,10,15,tzinfo=timezone.utc)
    return p

def slide_new(p):
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE);return s

def notes(s,body):
    s.notes_slide.notes_text_frame.text=('空白模板，不含真实人员、分值或业务资料。\n'+body+
    '\n成稿核验：所有占位已替换；数据来源和页码可追溯；行为影响标为假设；正文可编辑；全页渲染检查。\n保密：不要在共享备注中放真实姓名映射、测试账号、联系方式或无关内部资料。')

def header(s,num,title,subtitle):
    rect(s,f'p{num}.header.blue',.5,.38,.055,.34,BLUE)
    text(s,f'p{num}.title',title,.72,.23,11.65,.6,24,bold=True)
    text(s,f'p{num}.subtitle',subtitle,.72,.86,11.7,.34,11,MUTED)
    ln(s,f'p{num}.toprule',.53,1.28,12.81,1.28)
    text(s,f'p{num}.index',f'{num:02}',12.39,.32,.43,.36,12,BLUE,True,PP_ALIGN.RIGHT)
    ln(s,f'p{num}.footrule',.53,7.12,12.81,7.12)
    text(s,f'p{num}.footer','{批次} · 样本 {N} 人  |  内部使用 · 结论需结合实际表现验证',.55,7.21,9.0,.16,8.2,MUTED)
    text(s,f'p{num}.template_mark','空白模板 · 无真实数据',10.18,7.21,2.64,.16,8.2,MUTED,align=PP_ALIGN.RIGHT)

def block(s,name,x,y,w,h,title,body,evidence=None,color=BLUE):
    rect(s,name+'.bg',x,y,w,h,PALE if color==BLUE else ROSE,rounded=True)
    rect(s,name+'.stripe',x+.12,y+.19,.035,.24,color)
    text(s,name+'.title',title,x+.24,y+.09,w-.39,.39,14.0,color,True)
    text(s,name+'.body',body,x+.24,y+.53,w-.46,h-.72 if evidence is None else h-.98,12.1)
    if evidence:text(s,name+'.evidence',evidence,x+.24,y+h-.29,w-.45,.2,9.1,MUTED)

def radar(s,name,x,y,w,h,dims,caption,small=False):
    # Data-free guide: axes/labels and center placeholder only; replace the region with approved chart.
    n=len(dims);cx=x+w/2;cy=y+h*.51;r=min(w*.355,h*.345)
    f=7.5 if small else 12
    for q in (.2,.4,.6,.8,1):ring(s,f'{name}.ring{q}',cx-r*q,cy-r*q,2*r*q)
    for i,d in enumerate(dims):
        a=-math.pi/2+2*math.pi*i/n
        px=cx+r*math.cos(a);py=cy+r*math.sin(a)
        ln(s,f'{name}.axis{i}',cx,cy,px,py)
        rr=r*1.21;tx=cx+rr*math.cos(a);ty=cy+rr*math.sin(a)
        lw=.73 if small else 1.2;lh=.29 if small else .45
        label=d
        if small and len(d)>3:label=d[:2]+'\n'+d[2:]
        text(s,f'{name}.dim{i}',label,tx-lw/2,ty-lh/2,lw,lh,f,MUTED,align=PP_ALIGN.CENTER)
    rect(s,name+'.caption.bg',cx-w*.30,cy-.31,w*.60,.62,WHITE,LINE,True)
    text(s,name+'.replace',caption,cx-w*.29,cy-.29,w*.58,.58,7.8 if small else 13,MUTED,align=PP_ALIGN.CENTER)



def line_grid(s,name,x,y,w,h,dims):
    gx=x+.28;gy=y+.32;gw=w-.4;gh=h-.91
    for j in range(4):rect(s,f'{name}.band{j}',gx,gy+j*gh/4,gw,gh/4,'F2F6FD' if j%2==0 else 'F9FAFD')
    for j,val in enumerate((100,75,50,25,0)):
        yy=gy+j*gh/4;ln(s,f'{name}.grid{j}',gx,yy,gx+gw,yy,LINE,.75,True)
        text(s,f'{name}.tick{j}',str(val),x-.02,yy-.1,.26,.2,8,MUTED,align=PP_ALIGN.RIGHT)
    for i,d in enumerate(dims):
        xx=gx+(i+.5)*gw/len(dims);lab=d
        if len(d)>2:lab=d[:2]+'\n'+d[2:]
        text(s,f'{name}.dim{i}',lab,xx-gw/len(dims)/2,gy+gh+.1,gw/len(dims),.43,8.2,MUTED,align=PP_ALIGN.CENTER)
    text(s,name+'.replace','[按本批数据填入均值 / 中位值折线]',gx+.06,gy+gh*.43,gw-.12,.48,11,MUTED,align=PP_ALIGN.CENTER)


def personal():
    p=prs_new('Hogan个人报告空白模板');s=slide_new(p)
    text(s,'person.alias','{成员代号}',.408,.16,4.6,.46,25,BLUE,True)
    text(s,'person.batch','{届别 / 批次}  |  个人Hogan报告解读',7.1,.23,5.5,.3,10,MUTED,align=PP_ALIGN.RIGHT)
    ln(s,'header.rule.blue',.41,.658,6.66,.658,BLUE,1.1);ln(s,'header.rule.red',6.66,.658,12.91,.658,RED,1.1)
    rect(s,'core.bg',.383,.84,12.567,.56,PALE,rounded=True)
    text(s,'core.label','核心画像',.55,.98,.94,.22,11,BLUE,True)
    text(s,'core.text','[以HPI写清通常怎样做事、如何与人合作。]\n[补充最有解释力的HPI组合，不并排罗列高分项。]',1.70,.895,11.05,.44,13.6,bold=True)
    rect(s,'focus.bg',.383,1.47,12.567,.56,ROSE,rounded=True)
    text(s,'focus.label','需关注',.55,1.61,.94,.22,11,RED,True)
    text(s,'focus.text','[在具体工作场景中，可能出现什么行为、影响哪一步工作。]\n[分别说明重点观察什么，不写成已经发生的问题。]',1.70,1.515,11.05,.45,13.0)
    cols=[(.383,'HPI','日常工作风格',BLUE,['突出特点','群体差异','低位与关注']),
          (4.639,'HDS','压力下的行为风险',RED,['突出倾向','群体差异','风险与关注']),
          (8.894,'MVPI','动机与价值偏好',BLUE,['核心驱动','群体差异','低位与关注'])]
    boxes=[(1.30,2.552,2.442,2.24),(5.579,2.552,2.413,2.24),(9.877,2.552,2.327,2.24)]
    for (x,scale,sub,c,labs),b in zip(cols,boxes):
        rect(s,scale+'.panel',x,2.11,4.055,5.02,WHITE,LINE,True)
        rect(s,scale+'.header.bg',x+.015,2.123,4.025,.417,PALE if c==BLUE else ROSE)
        rect(s,scale+'.header.mark',x+.17,2.25,.035,.17,c)
        text(s,scale+'.header',scale,x+.28,2.17,.85,.29,15,c,True)
        text(s,scale+'.subtitle',sub,x+1.26,2.20,2.56,.25,10.2,INK,True)
        radar(s,scale+'.radar',*b,DIMS[scale],f'插入已确认的\n{scale}雷达图',True)
        text(s,scale+'.difference','最大差异：{维度}  {个人分} / 本批 {均值}',x+.19,4.835,3.69,.23,10.3,c,True)
        ln(s,scale+'.text.rule',x+.18,5.09,x+3.87,5.09)
        bodies={
         'HPI':['[单点或HPI组合体现的做事方式；\n不把风格写成能力。]','[比本批更突出或较少表现的倾向；\n低于均值不等于短板。]','[有依据的场景、行为与工作影响；\n没有低位，不硬凑缺点。]'],
         'HDS':['[何种压力情境下可能怎样反应；\n不要把量表名称当品行标签。]','[相对群体的压力反应差别；\n结合实际分档与分布。]','[对进度、交付或合作的影响；\n写明需要观察的行为。]'],
         'MVPI':['[哪些目标和工作环境更能激发投入；\n动机偏好不是能力。]','[本人与本批的投入来源差异；\n不是价值观优劣。]','[什么安排可能有效或不适合；\n应与本人确认，不预设事实。]']}
        for j,lab in enumerate(labs):
            yy=5.135+j*.663
            text(s,scale+f'.section{j}.label',lab,x+.19,yy,3.68,.19,9.6,c,True)
            text(s,scale+f'.section{j}.body',bodies[scale][j],x+.19,yy+.20,3.68,.43,11.9)
    text(s,'footer.method','群体基准：{本批N人，含本人}；分档：{核对当批报告}。行为与工作影响为待验证假设，不作为单一人事决策依据。',.41,7.22,11.7,.16,7.3,MUTED)
    text(s,'footer.internal','内部使用',12.08,7.22,.82,.16,7.5,MUTED,align=PP_ALIGN.RIGHT)
    notes(s,'核心画像：优先HPI单点与HPI内部组合；HDS/MVPI不固定塞入核心画像。\n原图替换区域以template-map.json的chart_regions为准，先移除对应radar.*占位对象，再等比放入确认图片。\n请填证据：来源文件/页码、原分/均值/差值、维度解释、组合假设、具体工作场景、观察重点。')
    return p


def group():
    p=prs_new('Hogan群体报告空白模板')
    s=slide_new(p)
    rect(s,'cover.blue',.8,1.20,.1,3.0,BLUE);rect(s,'cover.red',.8,4.31,.1,.53,RED)
    text(s,'cover.kicker','{届别 / 培养项目 / 批次}',1.14,1.2,10.7,.52,20,BLUE,True)
    text(s,'cover.title','Hogan\n群体分析报告',1.12,2.03,8.6,1.65,38,bold=True)
    text(s,'cover.subtitle','日常风格  ·  压力风险  ·  动机价值偏好',1.14,4.06,10.7,.39,18,MUTED)
    text(s,'cover.scope','样本 {N} 人  |  {分析目的 / 报告日期}',1.14,4.77,10.7,.35,14,MUTED)
    rect(s,'cover.readme.bg',1.13,5.63,11.06,.64,PALE,rounded=True)
    text(s,'cover.readme','看见群体共性与内部差异，具体说明对工作有哪些可能影响。',1.38,5.77,10.54,.29,15,BLUE)
    text(s,'cover.footer','内部使用  |  空白模板，不含真实测评信息',1.14,6.85,10.8,.25,10,MUTED)
    notes(s,'填写本批信息，不沿用历史年份/人数。不默认添加个人报告附录。')
    # 2
    s=slide_new(p);header(s,2,'整体画像：{先用一句话概括这批人的共同特点}','群体结论  /  先说行为与工作影响，再给少量数据')
    rect(s,'p2.summary.bg',.56,1.49,12.22,.55,PALE,rounded=True)
    text(s,'p2.summary','[两三种主要做事倾向怎样组合？对当前工作最值得关注的是什么？]',.78,1.59,11.77,.31,15,BLUE,True)
    for x,c,title in [(.58,BLUE,'可发挥的优势'),(6.88,RED,'重点关注的工作影响')]:
        text(s,f'p2.{title}.heading',title,x,2.18,5.72,.34,17,c,True)
        for j in range(4):
            yy=2.70+j*1.03
            if j%2:rect(s,f'p2.{title}.{j}.bg',x-.06,yy-.06,5.9,.95,LIGHT,rounded=True)
            text(s,f'p2.{title}.{j}.num',str(j+1),x+.05,yy,.28,.28,14,c,True)
            text(s,f'p2.{title}.{j}.title',f'[结论{j+1}：用完整的一句话说明]',x+.44,yy-.02,5.17,.28,13.4,INK,True)
            text(s,f'p2.{title}.{j}.body','[具体任务中可能怎样做，\n给进度、交付或合作带来什么影响。]',x+.44,yy+.32,5.17,.40,12)
            text(s,f'p2.{title}.{j}.evidence','依据：{关键分值 / 人数 / 分布}',x+.44,yy+.77,5.17,.18,8.8,MUTED)
    notes(s,'正文每条有工作行为和一步工作后果；群体/人数须原分支持，不仅凭均值推定多数。写法见writing-style.md。')
    # 3
    s=slide_new(p);header(s,3,'均值与中位值：{哪些接近，哪些还要看内部差异}','数据概览  /  同一批原分计算；统计不等于能力得分')
    for i,scale in enumerate(DIMS):
        x=.60+i*4.19;c=[BLUE,RED,GOLD][i]
        rect(s,f'p3.{scale}.header.bg',x,1.53,3.97,.86,[PALE,ROSE,'FAF7EE'][i],rounded=True)
        text(s,f'p3.{scale}.header',scale,x+.15,1.62,3.66,.26,17,c,True)
        text(s,f'p3.{scale}.summary','[该量表最主要的共性或分化]',x+.15,1.97,3.66,.25,11.6,MUTED)
        line_grid(s,f'p3.{scale}.chart',x,2.52,3.97,3.68,DIMS[scale])
        text(s,f'p3.{scale}.foot','[结合范围和分布解释，不用均值代替个人。]',x+.12,6.29,3.72,.44,11.4,MUTED)
    ln(s,'p3.legend.mean',8.55,6.84,8.98,6.84,BLUE,1.7);text(s,'p3.legend.mean.text','均值',9.05,6.7,.57,.27,10)
    ln(s,'p3.legend.median',10.16,6.84,10.59,6.84,RED,1.7,True);text(s,'p3.legend.median.text','中位值',10.66,6.7,.82,.27,10)
    notes(s,'折线区无真实数据。若需要仅均值可删中位值。坐标0–100，不将25/50/75网格当成HDS风险分档。均值与中位值接近不证明集中。')
    # 4
    s=slide_new(p);header(s,4,'HPI：{这批人平时怎样承担、学习、安排和合作}','日常工作风格  /  群体共性 + HPI内部组合 + 内部分化')
    radar(s,'p4.radar',.56,1.56,5.53,4.49,DIMS['HPI'],'插入本批HPI群体雷达图\n均值 / 中位值 / 个体轮廓')
    for j in range(3):
        block(s,f'p4.finding{j}',6.36,1.57+j*1.47,6.37,1.33,f'[结论{j+1}：具体的日常做事方式]',
              '[在常见任务中可能如何推进；\n说明对工作或合作哪一步有帮助或代价。]','依据：{少量数据；完整证据见受限备注}')
    rect(s,'p4.diff.bg',.66,6.24,12.05,.55,LIGHT,rounded=True)
    text(s,'p4.diff','内部差异：[哪些人或组合不能直接套用群体画像？默认用编号，不列真实姓名。]',.85,6.33,11.67,.35,12)
    notes(s,'HPI群体雷达图必须有统一范围、固定维度顺序和真实分值来源；解释以日常工作风格为主。不把高学习分等同学习能力。')
    # 5
    s=slide_new(p);header(s,5,'团队角色：{覆盖了什么，哪些工作需要主动安排}','角色覆盖  /  先确认项目公式；同一成员可兼有多个角色')
    radar(s,'p5.roles',.58,1.61,5.48,4.61,['结果','关系','过程','创新','务实'],'填入按认可规则统计的\n角色人数与占比')
    for j,(ti,bo) in enumerate([
       ('[优势：哪些自然倾向覆盖较多]','[写出在真实任务中可能怎样发挥作用，\n不要把角色人数等同实际能力。]'),
       ('[工作安排：哪些方面要形成配合]','[谁偏探索、谁重计划，怎样避免\n任务无人跟进或反复修改。]'),
       ('[边界：角色缺位不等于没有能力]','[未满足阈值只说明规则下未计入，\n不能据此判断有人不会做这类工作。]')]):
        block(s,f'p5.finding{j}',6.36,1.62+j*1.44,6.37,1.30,ti,bo,'依据：{人数 / N；公式见备注}')
    text(s,'p5.rules','规则：[填经确认的公式与来源]；占比=角色人数÷N，同一人可重复计入不同角色。无认可规则时不统计。',.72,6.53,11.95,.33,10,MUTED)
    notes(s,'用户确认既有项目规则后才使用：结果=抱负≥65；过程=审慎≥65；创新=好奇≥50且抱负≥35；务实=人际敏感度<65且好奇<50；关系=人际敏感度≥50且社交≥35。不是Hogan官方标准，不推断能力/绩效。')
    # 6
    s=slide_new(p);header(s,6,'HDS：{压力下哪些行为更值得观察}','压力倾向  /  群体主量表 → 可得子维度 → 场景与工作影响')
    radar(s,'p6.radar',.52,1.84,5.25,4.61,DIMS['HDS'],'插入本批HDS群体雷达图\n只按真实分值绘制')
    for j in range(4):
        x=6.06+(j%2)*3.40;y=1.60+(j//2)*2.47
        rect(s,f'p6.risk{j}.bg',x,y,3.24,2.27,ROSE,rounded=True)
        text(s,f'p6.risk{j}.title',f'{{重点量表 {j+1}}}',x+.17,y+.13,2.90,.33,16,RED,True)
        text(s,f'p6.risk{j}.counts','{达到关注区间的人数 / N}',x+.17,y+.56,2.90,.25,10.4,MUTED)
        text(s,f'p6.risk{j}.body','[什么情境下可能怎样做；\n会让哪一步工作受影响；\n具体要观察什么。]',x+.17,y+.96,2.90,.74,12.3)
        text(s,f'p6.risk{j}.subscale','子维度：{有图才统计；缺少则注明}',x+.17,y+1.91,2.90,.25,8.8,MUTED)
    text(s,'p6.limit','主量表分档与子维度“满格”分开。子维度缺页/看不清应注明；重点量表随本批数据选择。',.74,6.70,11.98,.27,10,MUTED)
    notes(s,'不得从主量表猜子维度格数；满4格是历史项目突出口径，并非主量表高风险等级；0人满格不证明没有相关倾向。')
    # 7
    s=slide_new(p);header(s,7,'HDS个体差异：{共性之外有哪些不同的压力反应}','编号热力图  /  不在正文或共享备注放姓名映射')
    x0,y0=.64,2.09;ww,hh=8.12,3.90;cw=ww/12;rh=hh/7
    labs=['编号']+DIMS['HDS']
    for rr in range(7):
        for cc in range(12):
            xx=x0+cc*cw;yy=y0+rr*rh
            rect(s,f'p7.cell{rr}.{cc}.bg',xx,yy,cw-.018,rh-.018,'EAF1FC' if rr==0 else ('F6F8FB' if rr%2 else WHITE),LINE)
            val=labs[cc] if rr==0 else (f'P{rr:02}' if cc==0 else '—')
            text(s,f'p7.cell{rr}.{cc}.text',val,xx+.015,yy+.035,cw-.046,rh-.087,9.3,BLUE if cc==0 else INK,rr==0,PP_ALIGN.CENTER)
    text(s,'p7.legend','图例：[根据本批已确认的HDS区间填写标签与颜色]',.67,1.62,8.13,.29,11.5,MUTED)
    for j in range(3):
        block(s,f'p7.pattern{j}',9.05,1.67+j*1.49,3.68,1.34,f'[行为组合 {j+1}]','[写共性或不同组合及工作影响，\n不逐人罗列负面标签。]',None,RED)
    text(s,'p7.expand','按实际N增减行或分页，不能靠缩小字号塞人；“—”是空白占位，不是0分。',.69,6.29,8.1,.46,11,MUTED)
    notes(s,'默认不写真实姓名/缩写映射；需要身份对照时单独受限保存。不得用多维高分总数进行人才排序。没有个人原分不能生成此页。')
    # 8
    s=slide_new(p);header(s,8,'MVPI：{什么更能激发投入，哪些偏好不同}','动机与价值偏好  /  写任务与环境的匹配，不写能力高低')
    radar(s,'p8.radar',.55,1.71,5.45,4.83,DIMS['MVPI'],'插入本批MVPI群体雷达图\n共同驱动与个体差异')
    for j in range(4):
        block(s,f'p8.finding{j}',6.31,1.60+j*1.23,6.42,1.09,f'[结论 {j+1}：更重视什么]',
            '[怎样的任务或工作安排可能更有吸引力，\n会影响投入的哪一方面；需与本人确认。]',None,BLUE)
    text(s,'p8.limit','传统低不等于不守规则；商业高不等于业务能力强。激励差异须逐人确认，不能只看平均分。',.74,6.70,11.95,.28,10,MUTED)
    notes(s,'MVPI解释工作价值偏好；材料不足时不声称最有效的激励方案，不做价值观优劣判断。')
    # 9
    s=slide_new(p);header(s,9,'群体总结：{最重要的工作影响和观察重点}','从测评倾向到实际任务  /  不只罗列分数，不使用抽象口号')
    xs=[.64,3.25,9.02];ws=[2.48,5.64,3.67]
    for j,lab in enumerate(['具体工作场景','可能的做法与工作影响','重点观察什么']):
        rect(s,f'p9.heading{j}.bg',xs[j],1.64,ws[j],.51,PALE)
        text(s,f'p9.heading{j}.text',lab,xs[j]+.14,1.74,ws[j]-.28,.28,14,BLUE,True)
    content=[('[日常推进任务]','[哪些倾向有助推进？什么情况下可能返工、\n遗漏跟进或拖慢关键工作？]','[谁负责、何时交、\n哪些地方要改做法？]'),
             ('[压力与分歧]','[被质疑、任务变动或时间紧时，可能怎样回应？\n对协作方、决策或交付会有什么影响？]','[是否听清不同意见，\n及时说明风险和困难？]'),
             ('[投入与工作安排]','[怎样的任务、反馈和环境更可能让人愿意投入？\n哪些差异需要分别沟通？]','[实际参与是否增加？\n任务结果是否更稳定？]')]
    for i,row in enumerate(content):
        yy=2.33+i*1.27
        for j,val in enumerate(row):
            rect(s,f'p9.row{i}.col{j}.bg',xs[j],yy,ws[j],1.08,LIGHT if i%2==0 else WHITE,LINE)
            text(s,f'p9.row{i}.col{j}.text',val,xs[j]+.16,yy+.13,ws[j]-.32,.80,12.4,INK,j==0)
    rect(s,'p9.summary.bg',.65,6.36,12.07,.40,ROSE,rounded=True)
    text(s,'p9.summary','[一句话总结：当前最要看清的工作行为是什么；不要直接评价潜力或绩效。]',.85,6.42,11.68,.26,12.4,RED,True)
    notes(s,'用户只要结论时不另加培训建议。发展建议若被允许，也须具体到任务安排、反馈动作和观察，而非笼统赋能。')
    # 10
    s=slide_new(p);header(s,10,'阅读指南：先看倾向与证据，再用工作表现验证','方法与保密边界  /  不以单项分数替代人才判断')
    defs=[('HPI','平时怎样做事','日常如何承担任务、安排工作、\n学习新东西和与人合作。',BLUE),
          ('HDS','压力下可能怎样反应','特定状态下某些倾向可能被放大，\n并影响任务推进或合作。',RED),
          ('MVPI','重视什么、为什么投入','哪些目标和工作环境更有吸引力；\n动机偏好不是能力。',GOLD)]
    for i,(name,sub,body,c) in enumerate(defs):
        x=.64+i*4.18
        rect(s,f'p10.{name}.bg',x,1.56,3.91,1.63,[PALE,ROSE,'FAF7EE'][i],rounded=True)
        text(s,f'p10.{name}.name',name,x+.18,1.73,3.55,.31,20,c,True)
        text(s,f'p10.{name}.sub',sub,x+.18,2.17,3.55,.25,13,INK,True)
        text(s,f'p10.{name}.body',body,x+.18,2.54,3.55,.44,11.6)
    for j,(title,body) in enumerate([
       ('统计与来源','样本：{N}；常模/版本：{来源原文}。\n均值、中位值、范围、人数均来自同一批可比数据。\n分档：{与当批报告核对后填写}。\n来源文件与实际页码：{仅写必要信息}。'),
       ('解释与保密','事实、计算、推断分开；没有业务背景就不编造案例。\n不把相对低分写成短板，不把风险倾向写成真实问题。\n群体报告默认编号，身份映射不进共享备注。\n真实数据、报告和图像不进入公共GitHub仓库。')]):
        x=.65+j*6.30
        text(s,f'p10.section{j}.title',title,x,3.61,5.86,.35,18,BLUE if j==0 else RED,True)
        text(s,f'p10.section{j}.body',body,x+.03,4.15,5.88,1.81,13.2)
    text(s,'p10.limit','测评用于发展性反馈；不是临床诊断，也不作为自动招聘筛选、排名、任用或淘汰依据。',.75,6.52,11.83,.37,12.2,MUTED)
    notes(s,'按实际输入填写，而非继承模板中的“已核对”。引用来源与计算方法写受限备注，身份映射始终单独保存。')
    return p


def clean_package(path:Path):
    # Remove default thumbnail; no inherited document preview or external relationships.
    out=path.with_suffix('.clean.pptx')
    with zipfile.ZipFile(path) as zin,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename.startswith('docProps/thumbnail.'):continue
            raw=zin.read(info.filename)
            if info.filename=='_rels/.rels':
                ET.register_namespace('', 'http://schemas.openxmlformats.org/package/2006/relationships')
                rt=ET.fromstring(raw)
                for el in list(rt):
                    if 'thumbnail' in el.get('Type',''):rt.remove(el)
                raw=ET.tostring(rt,encoding='utf-8',xml_declaration=True)
            zout.writestr(info.filename,raw)
    out.replace(path)


def build(out:Path)->dict:
    out.mkdir(parents=True,exist_ok=True);result={}
    for filename,creator in [('individual-report-template.pptx',personal),('group-report-template.pptx',group)]:
        p=creator();path=out/filename
        for sl in p.slides:
            for sh in sl.shapes:
                for el in sh._element.xpath('.//a:effectRef'):
                    el.set('idx','0')
                for sp in sh._element.xpath('./p:spPr'):
                    for old in sp.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst'):
                        sp.remove(old)
                    sp.append(OxmlElement('a:effectLst'))
        p.save(path);clean_package(path)
        pages=[]
        for i,sl in enumerate(p.slides,1):
            fields=[]
            for sh in sl.shapes:
                if sh.has_text_frame:
                    fields.append({'name':sh.name,'shape_id':sh.shape_id,'x':round(sh.left/914400,4),'y':round(sh.top/914400,4),
                                   'w':round(sh.width/914400,4),'h':round(sh.height/914400,4),'placeholder':sh.text})
            pages.append({'page':i,'fields':fields})
        result[filename]={'size_inches':[W,H],'page_count':len(p.slides),'pages':pages}
    result['individual-report-template.pptx']['chart_regions']={
        'HPI':{'x':1.30,'y':2.552,'w':2.442,'h':2.24,'remove_prefix':'HPI.radar.'},
        'HDS':{'x':5.579,'y':2.552,'w':2.413,'h':2.24,'remove_prefix':'HDS.radar.'},
        'MVPI':{'x':9.877,'y':2.552,'w':2.327,'h':2.24,'remove_prefix':'MVPI.radar.'}}
    result['instructions']='坐标为英寸；先移除对应图表占位形状，再等比插入核对通过的图；不要整页栅格化。名称用于定位，模板内无真实分值。'
    (out/'template-map.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return result


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--out',type=Path,default=Path('assets/templates'));a=ap.parse_args()
    if any((a.out/n).exists() for n in ('individual-report-template.pptx','group-report-template.pptx')):
        raise SystemExit('输出模板已经存在；请指定新目录，不覆盖已审核版本。')
    build(a.out);print('已生成1页个人模板、10页群体模板及字段映射（全部为空白结构）。')
if __name__=='__main__':main()
